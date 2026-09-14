import html
import re
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BODY_FONT = "Noto Sans CJK SC"
SAMPLE_LABEL = "【原剧第1–3集样例】"
NUMBER = r"(?:\d+(?:\.\d+)?|\.\d+)"
TIMING_SUFFIX = re.compile(rf"(?:[｜|· \t]*(?:预计|总时长|时长)[：:]?\s*{NUMBER}\s*(?:秒|s)(?:\s*[（(]\s*\d+\s*分\s*\d+\s*秒\s*[）)])?|[（(]\s*(?:(?:预计|总时长|时长)[：:]?\s*)?{NUMBER}\s*(?:秒|s)\s*[）)])$", re.IGNORECASE)


def episode_seconds(rows):
    total = Decimal(0)
    for row in rows:
        value = re.fullmatch(rf"({NUMBER})\s*(?:s|秒)?", row[6], re.IGNORECASE)
        if value is None or Decimal(value.group(1)) <= 0:
            raise ValueError(f"invalid duration for {row[0]}: {row[6]}")
        total += Decimal(value.group(1))
    return total.quantize(Decimal("0.000000001"), rounding=ROUND_HALF_UP)


def seconds_text(value):
    return format(value, "f").rstrip("0").rstrip(".") if "." in format(value, "f") else str(value)


def timing_summary(parsed):
    totals = [episode_seconds(rows) for _, _, rows in parsed]
    total = sum(totals)
    average = (total / len(totals)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    shots = sum(len(rows) for _, _, rows in parsed)
    return f"交付统计：共 {len(parsed)} 集、{shots} 镜；镜头合计 {seconds_text(total)} 秒；单集范围 {seconds_text(min(totals))}–{seconds_text(max(totals))} 秒，平均 {seconds_text(average)} 秒。"


def cells(line: str):
    return [item.strip() for item in line.strip().split("|")[1:-1]]


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(inches * 1440)))
    width.set(qn("w:type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def put_text(cell, text, *, bold=False, size=8, color=None, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text.replace("<br>", "\n"))
    run.bold = bold
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def parse_markdown(markdown: str):
    lines = [line for line in markdown.splitlines() if line.strip()]
    header_at = next(index for index, line in enumerate(lines) if line.startswith("|") and "镜头号" in line)
    header = cells(lines[header_at])
    rows = [cells(line) for line in lines[header_at + 2 :] if line.startswith("|")]
    headings = [line.lstrip("# ").strip() for line in lines[:header_at] if line.startswith("#")]
    heading = next((title for title in headings if not title.startswith(SAMPLE_LABEL)), "分镜剧本")
    shot = re.fullmatch(r"ep(\d+)-s\d+", rows[0][0], re.IGNORECASE) if rows else None
    for title in headings:
        numbered = re.match(r"^第\s*(\d+)\s*集", title)
        if numbered and (shot is None or int(numbered.group(1)) == int(shot.group(1))):
            heading = title
            break
    return heading, header, rows


def clean_episode_name(suffix):
    name = suffix.lstrip("｜|·:： \t").rstrip()
    while True:
        cleaned = TIMING_SUFFIX.sub("", re.sub(r"[｜|· \t]+(?:分镜表|分镜剧本|剧本)$", "", name)).rstrip()
        if cleaned == name:
            return name
        name = cleaned


def episode_heading(heading, rows, fallback_episode, screenplay_dir=None):
    shot = re.fullmatch(r"ep(\d+)-s\d+", rows[0][0], re.IGNORECASE) if rows else None
    existing = re.match(r"^第\s*(\d+)\s*集", heading)
    episode = int(shot.group(1)) if shot else int(existing.group(1)) if existing else fallback_episode
    name = clean_episode_name(heading[existing.end():] if existing else heading)
    if screenplay_dir is not None and name in ("", "分镜表", "分镜剧本", "剧本"):
        screenplay = Path(screenplay_dir) / f"ep-{episode:02d}.md"
        if screenplay.is_file():
            approved_heading = re.search(r"^#+[ \t]*第[ \t]*(\d+)[ \t]*集([^\r\n]*)", screenplay.read_text(encoding="utf-8"), re.MULTILINE)
            if approved_heading and int(approved_heading.group(1)) == episode:
                name = clean_episode_name(approved_heading.group(2))
    if name in ("", "分镜表", "分镜剧本", "剧本"):
        name = ""
    return f"第 {episode} 集" + (f"｜{name}" if name else "") + f"｜预计 {seconds_text(episode_seconds(rows))} 秒"


def add_storyboard(doc, heading, header, rows, *, page_break_before=False):
    if len(header) != 7 or any(len(row) != 7 for row in rows):
        raise ValueError("expected a strict 7-column storyboard table")
    title = doc.add_paragraph()
    title.style = doc.styles["Heading 1"]
    title.paragraph_format.page_break_before = page_break_before
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(heading)
    run.bold = True
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    widths = [0.60, 2.30, 2.25, 1.35, 1.50, 1.95, 0.55]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    header_row = table.rows[0]
    repeat_header(header_row)
    for index, value in enumerate(header):
        cell = header_row.cells[index]
        set_width(cell, widths[index])
        set_cell_margins(cell)
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        put_text(cell, value, bold=True, size=7, color=(255, 255, 255), center=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_width(cell, widths[index])
            set_cell_margins(cell)
            if row_index % 2:
                shade(cell, "F3F6FA")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            put_text(cell, value, size=8, center=index in (0, 6))
def main(input_path: Path, output_path: Path, delivery_title=None, screenplay_dir=None):
    chunks = [chunk for chunk in input_path.read_text(encoding="utf-8").split("\n---\n") if chunk.strip()]
    preface = chunks[0].splitlines()
    sample_title = next((line.lstrip("# ").strip() for line in preface if line.startswith("#") and line.lstrip("# ").startswith(SAMPLE_LABEL)), None)
    scope_note = next((line.strip() for line in preface if line.startswith("交付范围：")), None) if sample_title else None
    if sample_title:
        delivery_title = delivery_title or sample_title
        if SAMPLE_LABEL not in delivery_title:
            delivery_title = SAMPLE_LABEL + delivery_title
    parsed = [parse_markdown(chunk) for chunk in chunks]
    if any(len(header) != 7 or any(len(row) != 7 for row in rows) for _, header, rows in parsed):
        raise ValueError("expected a strict 7-column storyboard table")
    doc = Document()
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn("w:pBdr"))):
        title_properties.remove(border)
    # Keep the existing Title appearance while exposing episodes in document outlines.
    episode_style = doc.styles["Heading 1"]
    episode_style.base_style = doc.styles["Title"]
    for tag in ("w:pPr", "w:rPr"):
        properties = episode_style.element.find(qn(tag))
        if properties is not None:
            episode_style.element.remove(properties)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    episode_style.element.get_or_add_pPr().append(outline)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.28)
    section.bottom_margin = Inches(0.28)
    section.left_margin = Inches(0.22)
    section.right_margin = Inches(0.22)
    if delivery_title:
        doc.core_properties.title = delivery_title
        paragraph = doc.add_paragraph()
        paragraph.style = doc.styles["Title"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(delivery_title)
        run.bold = True
        run.font.name = BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0, 0, 0)
    if scope_note:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(scope_note)
        run.font.name = BODY_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(timing_summary(parsed))
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)
    for index, (heading, header, rows) in enumerate(parsed):
        add_storyboard(doc, episode_heading(heading, rows, index + 1, screenplay_dir), header, rows, page_break_before=index > 0)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    main(Path(sys.argv[1]), Path(sys.argv[2]), sys.argv[3] if len(sys.argv) > 3 else None, Path(sys.argv[4]) if len(sys.argv) > 4 else None)
